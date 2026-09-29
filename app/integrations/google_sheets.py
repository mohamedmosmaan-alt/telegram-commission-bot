"""
Reads the two raw tables from Google: the agent/phone directory and the
commission report table. Both are tabs in ONE Google Sheet (recommended —
see README), read via the free Sheets API with a service account.

An optional fallback lets you point at raw .xlsx files on Drive instead
(USE_EXCEL_FALLBACK=true), for anyone who can't convert to native Sheets —
each raw file is expected to have the SAME columns as the sheet version.
"""
import io

import gspread
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import openpyxl

from app.config.settings import settings
from app.utils.logging_config import logger

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.readonly",
]

AGENTS_REQUIRED_COLUMNS = {"master", "name", "phone num"}
COMMISSION_REQUIRED_COLUMNS = {"master"}


class DataSourceError(Exception):
    """Raised for any problem reaching or parsing the source data (§15)."""


def _get_credentials() -> Credentials:
    if not settings.google_service_account_json:
        raise DataSourceError("GOOGLE_SERVICE_ACCOUNT_JSON is not configured.")
    return Credentials.from_service_account_file(
        settings.google_service_account_json, scopes=SCOPES
    )


def _rows_from_table(header: list[str], rows: list[list], required: set[str]) -> list[dict]:
    header_clean = [str(h).strip() for h in header]
    header_lower = [h.lower() for h in header_clean]
    missing = required - set(header_lower)
    if missing:
        raise DataSourceError(f"Sheet is missing required column(s): {', '.join(missing)}")

    records = []
    for raw_row in rows:
        if not any(str(cell).strip() for cell in raw_row if cell is not None):
            continue  # skip fully empty rows (§15)
        padded = list(raw_row) + [""] * (len(header_clean) - len(raw_row))
        record = {header_clean[i]: padded[i] for i in range(len(header_clean))}
        records.append(record)
    return records


def _fetch_worksheet_values(worksheet_name: str) -> list[list]:
    try:
        creds = _get_credentials()
        client = gspread.authorize(creds)
        sheet = client.open_by_key(settings.google_sheet_id)
        worksheet = sheet.worksheet(worksheet_name)
        return worksheet.get_all_values()
    except gspread.exceptions.WorksheetNotFound as e:
        raise DataSourceError(f"Worksheet '{worksheet_name}' not found in the Google Sheet.") from e
    except gspread.exceptions.APIError as e:
        raise DataSourceError(f"Google Sheets API error on '{worksheet_name}': {e}") from e
    except Exception as e:
        raise DataSourceError(f"Could not read worksheet '{worksheet_name}': {e}") from e


def fetch_agents_rows() -> list[dict]:
    if settings.use_excel_fallback:
        return _fetch_excel_rows(settings.google_drive_excel_file_id_agents, AGENTS_REQUIRED_COLUMNS)
    values = _fetch_worksheet_values(settings.google_sheet_agents_worksheet)
    if not values:
        raise DataSourceError("Agents worksheet appears to be empty.")
    header, *rows = values
    return _rows_from_table(header, rows, AGENTS_REQUIRED_COLUMNS)


def fetch_commission_rows() -> list[dict]:
    if settings.use_excel_fallback:
        return _fetch_excel_rows(settings.google_drive_excel_file_id_commission, COMMISSION_REQUIRED_COLUMNS)
    values = _fetch_worksheet_values(settings.google_sheet_commission_worksheet)
    if not values:
        raise DataSourceError("Commission worksheet appears to be empty.")
    header, *rows = values
    return _rows_from_table(header, rows, COMMISSION_REQUIRED_COLUMNS)


def _fetch_excel_rows(file_id: str, required: set[str]) -> list[dict]:
    if not file_id:
        raise DataSourceError("USE_EXCEL_FALLBACK is true but a Drive file ID is missing.")
    try:
        creds = _get_credentials()
        drive_service = build("drive", "v3", credentials=creds)
        request = drive_service.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buffer.seek(0)
    except Exception as e:
        raise DataSourceError(f"Could not download Excel file {file_id} from Drive: {e}") from e

    try:
        workbook = openpyxl.load_workbook(buffer, data_only=True, read_only=True)
        worksheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        values = list(worksheet.iter_rows(values_only=True))
    except Exception as e:
        raise DataSourceError(f"Downloaded file {file_id} is not a valid Excel workbook: {e}") from e

    if not values:
        raise DataSourceError(f"Excel file {file_id} appears to be empty.")
    header = [str(h) if h is not None else "" for h in values[0]]
    rows = [list(r) for r in values[1:]]
    return _rows_from_table(header, rows, required)
